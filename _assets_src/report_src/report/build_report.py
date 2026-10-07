#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Генератор отчёта по курсовой «Основы веб-технологий» (GM GYM), версия 2.

Запуск одной командой пересобирает весь .docx:
    python3 /home/claude/v2/report/build_report.py

Что делает:
  1. готовит изображения (обрезка длинных скриншотов, сводные картинки, схема сайта) в report/build/img;
  2. собирает документ на основе шаблона титульного листа (title_template.docx);
  3. делает три прохода (номера рисунков/таблиц -> оглавление без номеров страниц ->
     оглавление с номерами страниц, которые берутся из PDF, полученного через LibreOffice);
  4. сохраняет результат в /home/claude/v2/out/Отчёт_курсовая_GM_GYM_v2.docx.

Тексты разделов лежат в content.py, картинки готовит images.py, оформление и вспомогательные функции здесь.
Когда будет готов макет Figma: положите PNG в /home/claude/v2/report_assets/figma/
(home_desktop, bjj_desktop, thanks_desktop, home_tablet, home_mobile, ui_kit, grids, prototype)
и запустите скрипт снова: вместо жёлтых плейсхолдеров подставятся картинки.
"""
import copy
import os
import re
import shutil
import subprocess
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT,
                            WD_TAB_LEADER)
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

PATHS = dict(
    template='/home/claude/inputs/title_template.docx',
    assets='/home/claude/v2/report_assets',
    screens='/home/claude/v2/report_assets/screens',
    repo='/home/claude/v2/site',
    build=os.path.join(HERE, 'build'),
    out='/home/claude/v2/out/Отчёт_курсовая_GM_GYM_v2.docx',
)
TEXT_WIDTH_CM = 16.5  # 21 - 3 - 1.5
FONT = 'Times New Roman'
MONO = 'Courier New'


# ---------------------------------------------------------------- низкоуровневые помощники XML
def set_rfonts(rpr, name):
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for a in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:' + a), name)


def shade(el_pr, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    el_pr.append(shd)


def style_font(st, name=FONT, size=14, bold=None, italic=None):
    st.font.name = name
    st.font.size = Pt(size)
    if bold is not None:
        st.font.bold = bold
    if italic is not None:
        st.font.italic = italic
    set_rfonts(st.element.get_or_add_rPr(), name)
    st.font.color.rgb = RGBColor(0, 0, 0)


def add_pstyle(doc, name, base='Normal', size=14, align=None, first=0.0, left=0.0, line=1.5,
               before=0, after=0, bold=None, italic=None, font=FONT, keep_next=False):
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles[base]
    st.quick_style = True
    style_font(st, font, size, bold, italic)
    pf = st.paragraph_format
    pf.first_line_indent = Cm(first)
    pf.left_indent = Cm(left)
    pf.right_indent = Cm(0)
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = keep_next
    pf.widow_control = True
    if align is not None:
        pf.alignment = align
    return st


# ---------------------------------------------------------------- изображения
from images import prepare_images  # noqa: E402


# ---------------------------------------------------------------- класс отчёта
class Report:
    def __init__(self, labels=None, headings=None, pages=None):
        self.labels = labels or {}      # {'f:key': n, 't:key': n, 'l:key': n}
        self.prev_headings = headings or []
        self.pages = pages or {}
        self.new_labels = {}
        self.headings = []              # (level, text, bookmark)
        self.cnt = dict(f=0, t=0, l=0)
        self.used_sources = set()
        self.bm_id = 100
        self.img_dir = os.path.join(PATHS['build'], 'img')
        self.doc = Document(PATHS['template'])
        self._fix_title_page()
        self._setup_styles()
        self._setup_sections()

    # ---- титульный лист
    def _fix_title_page(self):
        P = self.doc.paragraphs

        def set_run(par, old, new):
            for r in par.runs:
                if old in r.text:
                    r.text = r.text.replace(old, new)
                    return
            raise RuntimeError('не найдено в титульном листе: ' + old)

        set_run(P[12], 'ТЕМАТИКА ВАШЕГО ПРОЕКТА', 'Разработка веб-сайта клуба единоборств GM GYM')
        for shd in P[12]._p.xpath('.//w:shd[@w:fill="FFFF00"]'):   # жёлтая отметка «заполнить» в шаблоне
            shd.getparent().remove(shd)
        set_run(P[17], '241-671', '251-671')
        set_run(P[18], 'Фамилия И.О.', 'Смирнов Д.А.')
        set_run(P[24], 'Темникова К.Н.', 'Темникова К.Н.')   # без изменений, для проверки наличия
        set_run(P[34], '25', '26')
        self.title_last_par = P[34]
        cp = self.doc.core_properties
        cp.title = 'Отчёт по курсовой работе: Разработка веб-сайта клуба единоборств GM GYM'
        cp.author = 'Смирнов Даниил Алексеевич'
        cp.subject = 'Основы веб-технологий'
        cp.last_modified_by = 'Смирнов Даниил Алексеевич'

    # ---- стили
    def _setup_styles(self):
        doc = self.doc
        # отключаем автопереносы (в шаблоне включены)
        st = doc.settings.element
        for el in st.findall(qn('w:autoHyphenation')):
            st.remove(el)

        # встроенные Heading: переименовываем в канонические имена и оформляем
        by_id = {s.style_id: s for s in doc.styles}
        for sid, nm, size, before, after, lvl in (('1', 'heading 1', 16, 0, 12, 0),
                                                  ('2', 'heading 2', 14, 14, 6, 1),
                                                  ('3', 'heading 3', 14, 10, 4, 2)):
            s = by_id[sid]
            s.element.find(qn('w:name')).set(qn('w:val'), nm)
            style_font(s, FONT, size, True, False)
            if lvl == 2:
                s.font.italic = True
            pf = s.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.first_line_indent = Cm(0)
            pf.left_indent = Cm(0)
            pf.space_before = Pt(before)
            pf.space_after = Pt(after)
            pf.line_spacing = 1.15
            pf.keep_with_next = True
            pf.keep_together = True
            pf.page_break_before = (lvl == 0)
            ppr = s.element.get_or_add_pPr()
            for o in ppr.findall(qn('w:outlineLvl')):
                ppr.remove(o)
            ol = OxmlElement('w:outlineLvl')
            ol.set(qn('w:val'), str(lvl))
            ppr.append(ol)

        A = WD_ALIGN_PARAGRAPH
        add_pstyle(doc, 'Report Body', align=A.JUSTIFY, first=1.25, line=1.5)
        add_pstyle(doc, 'Report Plain', align=A.LEFT, first=0, line=1.5)
        add_pstyle(doc, 'Report Figure', size=12, align=A.CENTER, line=1.0, before=6, after=3, keep_next=True)
        add_pstyle(doc, 'Report Caption', size=12, align=A.CENTER, line=1.0, before=0, after=8, keep_next=False)
        add_pstyle(doc, 'Report Table Caption', size=12, align=A.LEFT, line=1.0, before=8, after=4, keep_next=True)
        add_pstyle(doc, 'Report Cell', size=11, align=A.LEFT, line=1.0, before=0, after=0)
        add_pstyle(doc, 'Report Code', size=10, align=A.LEFT, line=1.0, font=MONO)
        add_pstyle(doc, 'Report Spacer', size=6, align=A.LEFT, line=1.0)
        add_pstyle(doc, 'Report Placeholder', size=12, align=A.CENTER, line=1.0, before=6, after=3,
                   bold=True, keep_next=True)
        add_pstyle(doc, 'Report Note', size=12, align=A.LEFT, line=1.15, before=0, after=6, italic=True)
        add_pstyle(doc, 'Report Source', size=12, align=A.LEFT, first=0, left=0.9, line=1.0, after=4)
        add_pstyle(doc, 'Report Toc Title', align=A.CENTER, line=1.15, after=12, bold=True, size=16)
        add_pstyle(doc, 'Report Footer', size=12, align=A.CENTER, line=1.0)
        # стили оглавления (встроенные имена toc 1..3)
        for i, (ind, bold) in enumerate(((0, True), (0.6, False), (1.2, False)), start=1):
            s = add_pstyle(doc, 'toc %d' % i, size=12, align=A.LEFT, line=1.1, left=ind, after=2, bold=bold)
            s.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT,
                                                      WD_TAB_LEADER.DOTS)

    # ---- разделы и поля страницы
    def _setup_sections(self):
        doc = self.doc
        body = doc.element.body
        body_sect = body.find(qn('w:sectPr'))
        # раздел 1 (титульный лист): копия исходных свойств шаблона, без номера страницы
        s1 = copy.deepcopy(body_sect)
        tp = OxmlElement('w:titlePg')
        s1.append(tp)
        ppr = self.title_last_par._p.get_or_add_pPr()
        ppr.append(s1)
        # раздел 2: поля по заданию, нумерация продолжается
        sec = doc.sections[-1]
        sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
        sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
        sec.footer_distance = Cm(1.0)
        sec.header_distance = Cm(1.0)
        pg = body_sect.find(qn('w:pgNumType'))
        if pg is not None:
            body_sect.remove(pg)
        sec.footer.is_linked_to_previous = False
        fp = sec.footer.paragraphs[0]
        for r in list(fp.runs):
            r._r.getparent().remove(r._r)
        fp.style = doc.styles['Report Footer']
        self._field(fp, ' PAGE ', '2')

    def _field(self, par, instr, cached='1'):
        def r(child):
            run = OxmlElement('w:r')
            rpr = OxmlElement('w:rPr')
            set_rfonts(rpr, FONT)
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), '24')
            rpr.append(sz)
            run.append(rpr)
            run.append(child)
            par._p.append(run)

        b = OxmlElement('w:fldChar')
        b.set(qn('w:fldCharType'), 'begin')
        r(b)
        it = OxmlElement('w:instrText')
        it.set(qn('xml:space'), 'preserve')
        it.text = instr
        r(it)
        sp = OxmlElement('w:fldChar')
        sp.set(qn('w:fldCharType'), 'separate')
        r(sp)
        t = OxmlElement('w:t')
        t.text = cached
        r(t)
        e = OxmlElement('w:fldChar')
        e.set(qn('w:fldCharType'), 'end')
        r(e)

    # ---- ссылки и разметка текста
    def ref(self, s):
        def sub(m):
            kind, key = m.group(1), m.group(2)
            if kind == 's':
                nums = []
                for k in key.split(','):
                    k = k.strip()
                    if k not in self.src_index:
                        raise KeyError('нет источника: ' + k)
                    nums.append(str(self.src_index[k]))
                    self.used_sources.add(k)
                return ', '.join(nums)
            n = self.labels.get(kind + ':' + key)
            return str(n) if n is not None else '0'
        return re.sub(r'\{([ftl]|s):([^}]+)\}', sub, s)

    TOKEN = re.compile(r'(\*\*.+?\*\*|`.+?`|==.+?==|__.+?__)')

    def runs(self, par, text, size=None, bold=None, color=None):
        text = self.ref(text)
        for tok in self.TOKEN.split(text):
            if not tok:
                continue
            kw = {}
            if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
                tok, kw = tok[2:-2], dict(bold=True)
            elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
                tok, kw = tok[1:-1], dict(mono=True)
            elif tok.startswith('==') and tok.endswith('==') and len(tok) > 4:
                tok, kw = tok[2:-2], dict(hl=True)
            elif tok.startswith('__') and tok.endswith('__') and len(tok) > 4:
                tok, kw = tok[2:-2], dict(italic=True)
            run = par.add_run(tok)
            if kw.get('bold') or bold:
                run.bold = True
            if kw.get('italic'):
                run.italic = True
            if kw.get('mono'):
                run.font.name = MONO
                set_rfonts(run._r.get_or_add_rPr(), MONO)
                run.font.size = Pt((size or 14) - 1.5)
            elif size:
                run.font.size = Pt(size)
            if kw.get('hl'):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
        return par

    # ---- блоки
    def p(self, text, style='Report Body'):
        par = self.doc.add_paragraph(style=style)
        self.runs(par, text)
        return par

    def li(self, text):
        return self.p('– ' + text)

    def items(self, lst, numbered=False, end=';'):
        for i, t in enumerate(lst, 1):
            last = (i == len(lst))
            tail = '.' if last else end
            t = t.rstrip(';.')
            self.p((('%d) ' % i) if numbered else '– ') + t + tail)

    def note(self, text):
        return self.p(text, 'Report Note')

    def spacer(self):
        self.doc.add_paragraph(style='Report Spacer')

    def _heading(self, level, text, bookmark=True):
        par = self.doc.add_paragraph(style=self.doc.styles['heading %d' % level])
        name = '_Toc%06d' % (len(self.headings) + 1)
        self.bm_id += 1
        if bookmark:
            bs = OxmlElement('w:bookmarkStart')
            bs.set(qn('w:id'), str(self.bm_id))
            bs.set(qn('w:name'), name)
            par._p.append(bs)
        par.add_run(text)
        if bookmark:
            be = OxmlElement('w:bookmarkEnd')
            be.set(qn('w:id'), str(self.bm_id))
            par._p.append(be)
        self.headings.append((level, text, name))
        return par

    def h1(self, t): return self._heading(1, t)
    def h2(self, t): return self._heading(2, t)
    def h3(self, t): return self._heading(3, t)

    def _label(self, kind, key):
        self.cnt[kind] += 1
        if key:
            self.new_labels['%s:%s' % (kind, key)] = self.cnt[kind]
        return self.cnt[kind]

    # ---- рисунок
    def fig(self, name, caption, key=None, width=16.0, max_h=19.0, comment=None):
        path = name if os.path.isabs(name) else os.path.join(self.img_dir, name)
        if not os.path.exists(path):
            path = os.path.join(PATHS['screens'], name)
        with Image.open(path) as im:
            w, h = im.size
        wcm = width
        if wcm * h / w > max_h:
            wcm = max_h * w / h
        par = self.doc.add_paragraph(style='Report Figure')
        par.add_run().add_picture(path, width=Cm(wcm))
        n = self._label('f', key)
        cap = self.doc.add_paragraph(style='Report Caption')
        self.runs(cap, 'Рисунок %d — %s' % (n, caption))
        if comment:
            self.p(comment)

    def placeholder(self, what, hint=None, caption=None, key=None, comment=None, kind='СКРИНШОТ'):
        par = self.doc.add_paragraph(style='Report Placeholder')
        ppr = par._p.get_or_add_pPr()
        bd = parse_xml('<w:pBdr %s><w:top w:val="dashed" w:sz="8" w:space="6" w:color="7F7F7F"/>'
                       '<w:left w:val="dashed" w:sz="8" w:space="6" w:color="7F7F7F"/>'
                       '<w:bottom w:val="dashed" w:sz="8" w:space="6" w:color="7F7F7F"/>'
                       '<w:right w:val="dashed" w:sz="8" w:space="6" w:color="7F7F7F"/></w:pBdr>' % nsdecls('w'))
        ppr.append(bd)
        run = par.add_run()
        run.add_text('[ВСТАВИТЬ %s: %s]' % (kind, what))
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if hint:
            r2 = par.add_run()
            r2.add_break()
            r2.add_text(hint)
            r2.bold = False
            r2.font.size = Pt(11)
            r2.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if caption:
            n = self._label('f', key)
            cap = self.doc.add_paragraph(style='Report Caption')
            self.runs(cap, 'Рисунок %d — %s' % (n, caption))
        if comment:
            self.p(comment)

    def placeholder_text(self, text):
        """Одиночный жёлтый абзац (без рамки)."""
        par = self.doc.add_paragraph(style='Report Plain')
        r = par.add_run(text)
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        r.bold = True
        return par

    # ---- таблицы
    def _table_base(self, ncols, widths_cm):
        doc = self.doc
        t = doc.add_table(rows=0, cols=ncols)
        tblPr = t._tbl.tblPr
        for tag in ('w:tblW', 'w:tblLayout', 'w:tblBorders', 'w:tblCellMar', 'w:jc'):
            for e in tblPr.findall(qn(tag)):
                tblPr.remove(e)
        total = int(sum(widths_cm) * 567)
        tblPr.append(parse_xml('<w:tblW %s w:w="%d" w:type="dxa"/>' % (nsdecls('w'), total)))
        tblPr.append(parse_xml('<w:jc %s w:val="center"/>' % nsdecls('w')))
        tblPr.append(parse_xml(
            '<w:tblBorders %s>' % nsdecls('w') + ''.join(
                '<w:%s w:val="single" w:sz="4" w:space="0" w:color="000000"/>' % s
                for s in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')) + '</w:tblBorders>'))
        tblPr.append(parse_xml('<w:tblLayout %s w:type="fixed"/>' % nsdecls('w')))
        tblPr.append(parse_xml(
            '<w:tblCellMar %s><w:top w:w="40" w:type="dxa"/><w:left w:w="100" w:type="dxa"/>'
            '<w:bottom w:w="40" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tblCellMar>' % nsdecls('w')))
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn('w:gridCol')), widths_cm):
            gc.set(qn('w:w'), str(int(w * 567)))
        return t

    def _fill_cell(self, cell, spec, width_cm, size, align='l', bold=False, header=False):
        fill = None
        color = None
        if isinstance(spec, dict):
            text = spec.get('t', '')
            fill = spec.get('fill')
            color = spec.get('color')
            bold = spec.get('bold', bold)
            align = spec.get('align', align)
        else:
            text = spec
        tcPr = cell._tc.get_or_add_tcPr()
        for e in tcPr.findall(qn('w:tcW')):
            tcPr.remove(e)
        tcPr.insert(0, parse_xml('<w:tcW %s w:w="%d" w:type="dxa"/>' % (nsdecls('w'), int(width_cm * 567))))
        if fill:
            shade(tcPr, fill)
        va = OxmlElement('w:vAlign')
        va.set(qn('w:val'), 'center' if header else 'top')
        tcPr.append(va)
        lines = (text or '').split('\n')
        for i, ln in enumerate(lines):
            par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            par.style = self.doc.styles['Report Cell']
            par.alignment = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'c': WD_ALIGN_PARAGRAPH.CENTER,
                             'r': WD_ALIGN_PARAGRAPH.RIGHT}[align]
            self.runs(par, ln, size=size, bold=bold, color=color)

    def table(self, headers, rows, widths, caption=None, key=None, size=11, header_fill='D9D9D9',
              aligns=None, hdr_size=None, header_align='c'):
        if caption:
            n = self._label('t', key)
            cp = self.doc.add_paragraph(style='Report Table Caption')
            self.runs(cp, 'Таблица %d — %s' % (n, caption))
        t = self._table_base(len(widths), widths)
        aligns = aligns or ['l'] * len(widths)
        allrows = ([headers] if headers else []) + rows
        for ri, row in enumerate(allrows):
            cells = t.add_row().cells
            is_h = bool(headers) and ri == 0
            trPr = t.rows[-1]._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:cantSplit'))
            if is_h:
                trPr.append(OxmlElement('w:tblHeader'))
            for ci, spec in enumerate(row):
                if is_h:
                    spec = {'t': spec, 'fill': header_fill, 'bold': True, 'align': header_align} \
                        if not isinstance(spec, dict) else spec
                self._fill_cell(cells[ci], spec, widths[ci], hdr_size if (is_h and hdr_size) else size,
                                aligns[ci], header=is_h)
        self.spacer()
        return t

    def code(self, text, caption, key=None, lang=None):
        n = self._label('l', key)
        cp = self.doc.add_paragraph(style='Report Table Caption')
        self.runs(cp, 'Листинг %d — %s' % (n, caption))
        t = self._table_base(1, [TEXT_WIDTH_CM])
        row_ = t.add_row()
        row_._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        cell = row_.cells[0]
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.insert(0, parse_xml('<w:tcW %s w:w="%d" w:type="dxa"/>' % (nsdecls('w'), int(TEXT_WIDTH_CM * 567))))
        shade(tcPr, 'F2F2F2')
        lines = text.strip('\n').split('\n')
        for i, ln in enumerate(lines):
            if len(ln) > 78:
                print('  [!] длинная строка в листинге %s (%d): %s' % (key, len(ln), ln[:60]))
            par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            par.style = self.doc.styles['Report Code']
            run = par.add_run(ln)
            run.font.name = MONO
            set_rfonts(run._r.get_or_add_rPr(), MONO)
        self.spacer()

    # ---- оглавление
    def toc(self):
        doc = self.doc
        par = doc.add_paragraph(style='Report Toc Title')
        par.add_run('Оглавление')
        self.note('Оглавление собрано как поле Word (TOC \\o "1-3" \\h \\z \\u). Номера страниц в нём вычислены '
                  'при сборке отчёта; после правок откройте документ в Word, щёлкните по оглавлению правой кнопкой '
                  'мыши и выберите «Обновить поле» → «Обновить целиком» (или нажмите F9).')
        entries = self.prev_headings
        if not entries:
            entries = [(1, 'Оглавление будет сформировано при сборке', '_Toc000000')]
        for i, (lvl, text, bm) in enumerate(entries):
            tp = doc.add_paragraph(style='toc %d' % lvl)
            if i == 0:
                self._fld_char(tp, 'begin')
                it = OxmlElement('w:r')
                ins = OxmlElement('w:instrText')
                ins.set(qn('xml:space'), 'preserve')
                ins.text = ' TOC \\o "1-3" \\h \\z \\u '
                it.append(ins)
                tp._p.append(it)
                self._fld_char(tp, 'separate')
            pg = str(self.pages.get(bm, ''))
            hl = parse_xml('<w:hyperlink %s w:anchor="%s" w:history="1"/>' % (nsdecls('w'), bm))
            for txt, is_tab in ((text, False), (None, True), (pg, False)):
                r = OxmlElement('w:r')
                if is_tab:
                    r.append(OxmlElement('w:tab'))
                else:
                    t = OxmlElement('w:t')
                    t.set(qn('xml:space'), 'preserve')
                    t.text = txt
                    r.append(t)
                hl.append(r)
            tp._p.append(hl)
            if i == len(entries) - 1:
                self._fld_char(tp, 'end')

    def _fld_char(self, par, typ):
        r = OxmlElement('w:r')
        fc = OxmlElement('w:fldChar')
        fc.set(qn('w:fldCharType'), typ)
        r.append(fc)
        par._p.append(r)

    def save(self, path):
        self.doc.save(path)


# ---------------------------------------------------------------- проходы и PDF
def to_pdf(docx_path, outdir):
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', outdir, docx_path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    return os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + '.pdf')


def pdf_pages(pdf):
    txt = subprocess.run(['pdftotext', '-layout', pdf, '-'], check=True, capture_output=True).stdout.decode('utf-8')
    pages = txt.split('\f')
    if pages and not pages[-1].strip():
        pages.pop()
    return pages


def norm(s):
    return re.sub(r'\s+', ' ', s).strip()


def locate_headings(headings, pages_txt):
    """Находит страницу каждого заголовка (по порядку, номера страниц = индексы PDF + 1)."""
    npages = [norm(p) for p in pages_txt]
    # начало основной части: первая страница, начинающаяся с «Введение»
    start = next(i for i, p in enumerate(npages) if i >= 2 and p.startswith('Введение'))
    res, cur = {}, start
    for lvl, text, bm in headings:
        key = norm(text)
        for i in range(cur, len(npages)):
            if key in npages[i]:
                res[bm] = i + 1
                cur = i
                break
        else:
            raise RuntimeError('заголовок не найден в PDF: ' + text)
    return res


def build_once(labels, headings, pages):
    import content
    rep = Report(labels=labels, headings=headings, pages=pages)
    content.build(rep)
    return rep


def main():
    prepare_images()
    os.makedirs(PATHS['build'], exist_ok=True)
    tmp = os.path.join(PATHS['build'], 'pass.docx')
    # проход A: нумерация рисунков, таблиц и заголовки
    a = build_once({}, [], {})
    # проход B: оглавление без номеров страниц (чтобы длина оглавления была как в финале)
    b = build_once(a.new_labels, a.headings, {})
    b.save(tmp)
    pdf = to_pdf(tmp, PATHS['build'])
    pages_map = locate_headings(b.headings, pdf_pages(pdf))
    # проход C: финальный
    c = build_once(a.new_labels, b.headings, pages_map)
    os.makedirs(os.path.dirname(PATHS['out']), exist_ok=True)
    c.save(PATHS['out'])
    # проверка: страницы не сдвинулись после вставки номеров
    pdf2 = to_pdf(PATHS['out'], PATHS['build'])
    pages2 = locate_headings(c.headings, pdf_pages(pdf2))
    if pages2 != pages_map:
        # ещё один проход с уточнёнными номерами
        c = build_once(a.new_labels, c.headings, pages2)
        c.save(PATHS['out'])
        pdf2 = to_pdf(PATHS['out'], PATHS['build'])
    n = len(pdf_pages(pdf2))
    print('Готово:', PATHS['out'])
    print('Страниц (по LibreOffice):', n)
    print('Рисунков:', c.cnt['f'], ' таблиц:', c.cnt['t'], ' листингов:', c.cnt['l'])


if __name__ == '__main__':
    main()
